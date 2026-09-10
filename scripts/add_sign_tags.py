#!/usr/bin/env python3
"""Ensure a Doc Gen template carries valid Box Sign fields for recipient 1.

https://support.box.com/hc/en-us/articles/4404085855251-Creating-Templates-Using-Tags
"""
from pathlib import Path
import sys
from docx import Document
from docx.shared import Pt
try:
    from .generate_docgen_templates import add_signing_fields
except ImportError:
    from generate_docgen_templates import add_signing_fields


def add_sign_tags_to_template(template_path: str, output_path: str | None = None):
    doc = Document(template_path)
    replacements = {
        "[[s:email:signer]]": "[[s|1|id:borrower_signature                    ]]",
        "[[d:email:signer]]": "[[d|1|id:borrower_signed_date]]",
        "{{loan.borrowerEntity}}": "{{loan.borrower}}",
    }
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            for old, new in replacements.items():
                if old in run.text:
                    run.text = run.text.replace(old, new)
            if "[[s|1|id:borrower_signature" in run.text:
                run.font.size = Pt(18)
    text = "\n".join(p.text for p in doc.paragraphs)
    if "[[s|1" not in text:
        add_signing_fields(doc)
    doc.save(output_path or template_path)
    return True


if __name__ == "__main__":
    source = Path(sys.argv[1] if len(sys.argv) > 1 else "output/docgen/los-commitment-letter-template.docx")
    if not source.exists():
        raise SystemExit(f"Template not found: {source}")
    add_sign_tags_to_template(str(source), sys.argv[2] if len(sys.argv) > 2 else None)
