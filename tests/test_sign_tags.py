import tempfile
import unittest
from pathlib import Path
from docx import Document
from scripts.generate_docgen_templates import commitment_letter
from scripts.add_sign_tags import add_sign_tags_to_template


class SignTagTests(unittest.TestCase):
    def test_generated_template_assigns_signature_and_date_to_recipient_one(self):
        text = '\n'.join(p.text for p in commitment_letter().paragraphs)
        self.assertIn('[[s|1|id:borrower_signature', text)
        self.assertIn('[[d|1|id:borrower_signed_date]]', text)
        self.assertNotIn(':email:signer', text)
        tag = next(r for p in commitment_letter().paragraphs for r in p.runs if '[[s|1' in r.text)
        self.assertEqual(tag.font.size.pt, 18)

    def test_legacy_repair_is_idempotent_and_preserves_other_text(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'template.docx'
            doc = Document()
            doc.add_paragraph('Original loan terms')
            doc.add_paragraph('Borrower signature [[s:email:signer]]')
            doc.add_paragraph('Date [[d:email:signer]]')
            doc.add_paragraph('{{loan.borrowerEntity}}')
            doc.save(path)
            add_sign_tags_to_template(str(path))
            add_sign_tags_to_template(str(path))
            text = '\n'.join(p.text for p in Document(path).paragraphs)
            self.assertEqual(text.count('[[s|1'), 1)
            self.assertEqual(text.count('[[d|1'), 1)
            self.assertIn('Original loan terms', text)
            self.assertIn('{{loan.borrower}}', text)
            self.assertNotIn(':email:signer', text)
            tag = next(r for p in Document(path).paragraphs for r in p.runs if '[[s|1' in r.text)
            self.assertEqual(tag.font.size.pt, 18)

    def test_plain_template_gets_one_signature_block(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'template.docx'
            Document().save(path)
            add_sign_tags_to_template(str(path))
            add_sign_tags_to_template(str(path))
            self.assertEqual('\n'.join(p.text for p in Document(path).paragraphs).count('[[s|1'), 1)
